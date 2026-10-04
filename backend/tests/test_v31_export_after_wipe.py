"""Roadmap Nr. 5 (Forum #56, Bacardi): nach "Alle Daten loeschen" lieferten
PDF/Excel-Export noch alle Geraete, und die eigene Feldauswahl fiel auf
Standard zurueck.

* Export nach ``/devices/bulk/delete-all`` ist leer (Geraete liegen im
  Papierkorb, der Export liest nur ``deleted_at IS NULL``).
* Die Feldauswahl liegt seit v3.1.0 auf dem Server und ueberlebt das Loeschen.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


@pytest.fixture
def client(tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(settings, "PHOTOS_DIR", tmp_path / "photos")

    from fastapi import FastAPI
    from app.database import init_db
    from app.routers import devices, export, settings as settings_router

    init_db()
    app = FastAPI()
    app.include_router(devices.router, prefix="/api")
    app.include_router(export.router, prefix="/api")
    app.include_router(settings_router.router, prefix="/api")
    return TestClient(app)


def _insert(uuid: str) -> None:
    from app.database import get_db

    with get_db() as conn:
        conn.execute(
            "INSERT INTO devices (uuid, typ, bezeichnung) VALUES (?, 'Sensor', ?)",
            (uuid, f"Geraet {uuid}"),
        )


def _xlsx_names(content: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content))
    values: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values.extend(str(v) for v in row if v is not None)
    return values


def test_export_is_empty_after_delete_all(client):
    _insert("a1")
    _insert("a2")
    before = client.get("/api/export/xlsx", params={"fields": "bezeichnung"})
    assert any("Geraet a1" in v for v in _xlsx_names(before.content))

    r = client.post("/api/devices/bulk/delete-all", json={})
    assert r.status_code == 200 and r.json()["deleted"] == 2

    after = client.get("/api/export/xlsx", params={"fields": "bezeichnung"})
    assert after.status_code == 200
    assert not any("Geraet a" in v for v in _xlsx_names(after.content))
    pdf = client.get("/api/export/pdf", params={"fields": "bezeichnung"})
    assert pdf.status_code == 200


def test_export_field_selection_is_kept_on_the_server(client):
    assert client.get("/api/settings/export_fields").json() == {"fields": None, "preset": None}

    client.post("/api/settings/export_fields",
                json={"fields": ["nr", "bezeichnung", "standort_name"], "preset": None})
    _insert("b1")
    client.post("/api/devices/bulk/delete-all", json={})

    got = client.get("/api/settings/export_fields").json()
    assert got["fields"] == ["nr", "bezeichnung", "standort_name"]
    assert got["preset"] is None

    client.post("/api/settings/export_fields",
                json={"fields": ["nr", "  ", "typ"], "preset": "rueckbau"})
    got = client.get("/api/settings/export_fields").json()
    assert got == {"fields": ["nr", "typ"], "preset": "rueckbau"}
