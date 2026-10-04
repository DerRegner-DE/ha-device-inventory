"""v3.1.0 — Forum-Zusagen aus Thread 92177 #32 und Forum #90.

* Nr. 19: Feld "Wandschalter gebrueckt" (yes/no + Freitext), sprachneutral
  gespeichert, in device_history verfolgt, Pflichtspalte im Preset rueckbau.
* Nr. 20: Preset "notfallmappe" -- ohne Netzwerkdetails, ohne HA-Interna.
* Nr. 16: Export sortiert nach Etage > Standort > Bezeichnung, Excel dann als
  flache Tabelle ohne Kategorie-Zwischenzeilen; "Etage" zeigt den HA-Namen.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


@pytest.fixture
def client(tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(settings, "PHOTOS_DIR", tmp_path / "photos")

    from fastapi import FastAPI
    from app.database import init_db
    from app.routers import devices, export

    # MQTT-Veroeffentlichung im Hintergrund ist hier egal.
    async def _noop(*a, **kw):
        return None

    monkeypatch.setattr(devices, "publish_device", _noop)

    init_db()
    app = FastAPI()
    app.include_router(devices.router, prefix="/api")
    app.include_router(export.router, prefix="/api")
    return TestClient(app)


def test_bridged_switch_is_stored_language_neutrally_and_tracked(client):
    r = client.post("/api/devices", json={
        "typ": "Aktor/Relais", "bezeichnung": "Shelly Flur",
        "schalter_gebrueckt": "Yes ", "schalter_gebrueckt_hinweis": "Shelly detached, Taster am Eingang",
    })
    assert r.status_code == 201, r.text
    dev = r.json()
    assert dev["schalter_gebrueckt"] == "yes"

    r = client.put(f"/api/devices/{dev['uuid']}", json={"schalter_gebrueckt": "ja bitte"})
    assert r.json()["schalter_gebrueckt"] is None, "nur yes/no werden gespeichert"

    from app.services.history import TRACKED_FIELDS
    assert {"schalter_gebrueckt", "schalter_gebrueckt_hinweis"} <= TRACKED_FIELDS


def test_rueckbau_preset_requires_the_bridged_switch_column():
    from app.routers.export import EXPORT_FIELD_PRESETS

    assert "schalter_gebrueckt" in EXPORT_FIELD_PRESETS["rueckbau"]
    assert "schalter_gebrueckt_hinweis" in EXPORT_FIELD_PRESETS["rueckbau"]


def test_notfallmappe_preset_is_for_the_helper_not_the_admin():
    from app.routers.export import EXPORT_FIELD_PRESETS
    from app.services.excel_export import FIELD_LABELS
    from app.services.pdf_export import FIELD_LABELS_DE, FIELD_LABELS_EN

    fields = EXPORT_FIELD_PRESETS["notfallmappe"]
    for needed in ("bezeichnung", "standort_name", "hersteller", "modell", "seriennummer",
                   "ohne_ha", "schalter_gebrueckt", "external_url"):
        assert needed in fields
    for noise in ("ip_adresse", "mac_adresse", "firmware", "integration",
                  "ha_device_id", "ha_entity_id", "netzwerk"):
        assert noise not in fields
    for f in fields:
        assert f in FIELD_LABELS and f in FIELD_LABELS_DE and f in FIELD_LABELS_EN, f


def _rows(content: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(content)).active
    return [r for r in ws.iter_rows(values_only=True)]


def _seed() -> None:
    from app.database import get_db

    with get_db() as conn:
        conn.execute("INSERT INTO ha_areas (area_id, name, floor_id, floor_name) "
                     "VALUES ('kueche', 'Kueche', 'eg', 'Erdgeschoss')")
        conn.execute("INSERT INTO ha_areas (area_id, name, floor_id, floor_name) "
                     "VALUES ('bad', 'Bad', 'og', 'Obergeschoss')")
        for uuid, name, floor, place, integ in (
            ("1", "Zeta Licht", "og", "Bad", "zha"),
            ("2", "Alpha Licht", "og", "Bad", "fritz"),
            ("3", "Herd", "eg", "Kueche", "home_connect"),
            ("4", "Ohne Ort", None, None, "zha"),
        ):
            conn.execute(
                "INSERT INTO devices (uuid, typ, bezeichnung, standort_floor_id, standort_name, "
                "integration, schalter_gebrueckt) VALUES (?, 'Sensor', ?, ?, ?, ?, 'yes')",
                (uuid, name, floor, place, integ),
            )


def test_xlsx_sorted_by_location_is_flat_and_shows_floor_names(client):
    _seed()
    r = client.get("/api/export/xlsx", params={
        "fields": "standort_floor_id,standort_name,bezeichnung,schalter_gebrueckt",
        "sort": "standort",
    })
    rows = _rows(r.content)
    assert rows[0] == ("Etage", "Standort", "Bezeichnung", "Schalter gebrückt")
    assert [row[2] for row in rows[1:]] == ["Herd", "Alpha Licht", "Zeta Licht", "Ohne Ort"]
    assert rows[1][0] == "Erdgeschoss"
    assert rows[1][3] == "Ja"


def test_xlsx_default_keeps_category_groups(client):
    _seed()
    r = client.get("/api/export/xlsx", params={"fields": "nr,bezeichnung"})
    rows = _rows(r.content)
    # Gruppenzeilen haben in Spalte 1 nichts, in Spalte 2 den Kategorienamen.
    assert any(row[0] is None and row[1] for row in rows[1:])


def test_pdf_sorted_by_location_renders(client):
    _seed()
    from app.routers.export import EXPORT_FIELD_PRESETS

    r = client.get("/api/export/pdf", params={
        "fields": ",".join(EXPORT_FIELD_PRESETS["notfallmappe"]), "sort": "standort",
    })
    assert r.status_code == 200
    assert r.content.startswith(b"%PDF")
